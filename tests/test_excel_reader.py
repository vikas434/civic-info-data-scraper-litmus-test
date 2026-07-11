"""Tests for Excel workbook reader."""

from pathlib import Path

import openpyxl
import pytest

from app.excel_reader import read_zip_groups


def make_workbook(tmp_path: Path, rows: list[tuple]) -> Path:
    """Create a minimal test workbook with standard headers."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Zip Code", "Gov Level", "Branch", "Role", "MyCivX \u2013 Name", "MyCivX \u2013 Website"])
    for row in rows:
        ws.append(list(row))
    path = tmp_path / "test.xlsx"
    wb.save(path)
    return path


@pytest.fixture()
def simple_workbook(tmp_path: Path) -> Path:
    return make_workbook(tmp_path, [
        ("10019", "State (New York)", "Executive", "Governor", "Kathy Hochul", "https://www.ny.gov/governor"),
        (None,    "Local (NYC)",      "Executive", "Mayor of NYC", "Eric Adams", "https://www.nyc.gov/mayors-office"),
        ("10012", "State (New York)", "Executive", "Governor", "[To be filled]", None),
        (None,    "Local (NYC)",      "Executive", "Mayor of NYC", "[To be filled]", None),
    ])


def test_returns_correct_number_of_groups(simple_workbook: Path) -> None:
    groups, _ = read_zip_groups(simple_workbook)
    assert len(groups) == 2


def test_zip_codes_are_correct(simple_workbook: Path) -> None:
    groups, _ = read_zip_groups(simple_workbook)
    assert groups[0].zip_code == "10019"
    assert groups[1].zip_code == "10012"


def test_forward_fill_zip(simple_workbook: Path) -> None:
    groups, _ = read_zip_groups(simple_workbook)
    # Both rows under 10019 should be in the same group
    assert len(groups[0].roles) == 2


def test_role_names_are_correct(simple_workbook: Path) -> None:
    groups, _ = read_zip_groups(simple_workbook)
    assert groups[0].roles[0].name == "Governor"
    assert groups[0].roles[1].name == "Mayor of NYC"


def test_mycivx_name_is_read(simple_workbook: Path) -> None:
    groups, _ = read_zip_groups(simple_workbook)
    assert groups[0].roles[0].mycivx_name == "Kathy Hochul"


def test_mycivx_website_is_read(simple_workbook: Path) -> None:
    groups, _ = read_zip_groups(simple_workbook)
    assert groups[0].roles[0].mycivx_website == "https://www.ny.gov/governor"


def test_placeholder_becomes_none(simple_workbook: Path) -> None:
    groups, _ = read_zip_groups(simple_workbook)
    # 10012 has "[To be filled]" placeholders
    assert groups[1].roles[0].mycivx_name is None
    assert groups[1].roles[0].mycivx_website is None


def test_empty_mycivx_website_is_none(simple_workbook: Path) -> None:
    groups, _ = read_zip_groups(simple_workbook)
    assert groups[1].roles[0].mycivx_website is None


def test_row_indices_are_correct(simple_workbook: Path) -> None:
    groups, _ = read_zip_groups(simple_workbook)
    # Header is row 1, data starts row 2
    assert groups[0].row_indices[0] == 2
    assert groups[0].row_indices[1] == 3
    assert groups[1].row_indices[0] == 4


def test_gov_level_is_read(simple_workbook: Path) -> None:
    groups, _ = read_zip_groups(simple_workbook)
    assert groups[0].roles[0].gov_level == "State (New York)"


def test_branch_is_read(simple_workbook: Path) -> None:
    groups, _ = read_zip_groups(simple_workbook)
    assert groups[0].roles[0].branch == "Executive"


def test_emoji_in_header_is_handled(tmp_path: Path) -> None:
    """Column headers with emoji prefixes should still be detected."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([
        "Zip Code", "Gov Level", "Branch", "Role",
        "\U0001f310 MyCivX \u2013 Name", "\U0001f310 MyCivX \u2013 Website",
    ])
    ws.append(["10019", "Federal", "Executive", "President", "Donald Trump", "https://www.whitehouse.gov"])
    path = tmp_path / "emoji.xlsx"
    wb.save(path)

    groups, _ = read_zip_groups(path)
    assert groups[0].roles[0].mycivx_name == "Donald Trump"


def test_no_anomalies_for_clean_workbook(simple_workbook: Path) -> None:
    _, anomalies = read_zip_groups(simple_workbook)
    assert anomalies == []
